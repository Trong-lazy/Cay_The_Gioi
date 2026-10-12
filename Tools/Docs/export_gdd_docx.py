"""Export the authoritative Markdown GDD to a readable Word document."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "Docs" / "GDD_AI_Agent.md"
OUTPUT = ROOT / "Docs" / "GDD_AI_Agent.docx"


def clean(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("`", "").replace("**", "")


def table(doc: Document, lines: list[str]) -> None:
    rows = [[clean(c.strip()) for c in line.strip().strip("|").split("|")] for line in lines]
    rows.pop(1)
    t = doc.add_table(rows=1, cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    t.autofit = False
    # 6.5 inches usable width on Letter; give descriptive columns more room.
    widths = [1.35, 1.55, 1.5, 2.1] if len(rows[0]) == 4 else [3.25] * len(rows[0])
    if len(rows[0]) != 4:
        widths = [6.5 / len(rows[0])] * len(rows[0])
    for column, width in zip(t.columns, widths):
        column.width = Inches(width)
    for ri, row_text in enumerate(rows):
        row = t.rows[0] if ri == 0 else t.add_row()
        properties = row._tr.get_or_add_trPr()
        split = OxmlElement("w:cantSplit")
        properties.append(split)
        if ri == 0:
            repeat = OxmlElement("w:tblHeader")
            properties.append(repeat)
        for ci, (cell, text) in enumerate(zip(row.cells, row_text)):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                item = OxmlElement(f"w:{side}")
                item.set(qn("w:w"), "100")
                item.set(qn("w:type"), "dxa")
                margins.append(item)
            tcpr.append(margins)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "233F4D" if ri == 0 else ("F2F5F6" if ri % 2 == 0 else "FFFFFF"))
            tcpr.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(text)
            run.font.size = Pt(10.5)
            if ri == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def main() -> None:
    doc = Document()
    # The bundled default template carries an inherited blue Title border.
    # Remove paragraph borders from styles so the title stays plain and black.
    for element in list(doc.styles.element.iter(qn("w:pBdr"))):
        element.getparent().remove(element)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    sec.left_margin = sec.right_margin = Inches(1)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
        style.element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(5)
    title = doc.styles["Title"]
    title.font.size = Pt(24)
    title.paragraph_format.space_after = Pt(14)
    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        h = doc.styles[name]
        h.font.size = Pt(size)
        h.font.bold = True
        h.paragraph_format.space_before = Pt(13)
        h.paragraph_format.space_after = Pt(7)
        h.paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    r = footer.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    r._r.addnext(field)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            table(doc, block)
            continue
        if line.startswith("# "):
            doc.add_paragraph(clean(line[2:]), style="Title")
        elif line.startswith("## "):
            doc.add_paragraph(clean(line[3:]), style="Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(clean(line[4:]), style="Heading 2")
        elif line.startswith("- "):
            doc.add_paragraph(clean(line[2:]), style="List Bullet")
        elif re.match(r"\d+\. ", line):
            doc.add_paragraph(clean(re.sub(r"^\d+\. ", "", line)), style="List Number")
        else:
            doc.add_paragraph(clean(line))
        i += 1
    doc.core_properties.title = "GDD Cây Thế Giới dành cho AI Agent"
    doc.core_properties.subject = "Khung game và quy trình triển khai trong Unity"
    doc.core_properties.author = ""
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Exported {OUTPUT}")


if __name__ == "__main__":
    main()
